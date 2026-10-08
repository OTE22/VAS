"""Build the editable guide. Run with python-docx, Pillow and PyMuPDF installed."""
from pathlib import Path
from collections import Counter
import json, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
WS=json.loads((ROOT/'source/workflows.json').read_text()); COV=json.loads((ROOT/'source/coverage.json').read_text())
FIG=json.loads((ROOT/'source/figure-index.json').read_text())
PAGES=json.loads((ROOT/'source/page-map.json').read_text()) if (ROOT/'source/page-map.json').exists() else {}
D=Document();NAVY='17324B';TEAL='087F83';GRAY='526373';ORANGE='B65324'
sections={1:'Introduction, audience and prerequisites',2:'Login and role differences',3:'Interface orientation',4:'Quick-start workflow',5:'Everyday operator tasks',6:'Advanced and administrator tasks',7:'Troubleshooting and common mistakes',8:'Capability coverage and remaining gaps'}
styles=D.styles
for n in ['Normal','Body Text','List Bullet','List Number']:
 styles[n].font.name='Calibri';styles[n].font.size=Pt(11);styles[n].paragraph_format.space_after=Pt(7)
for n,sz in [('Title',34),('Heading 1',24),('Heading 2',18),('Heading 3',13)]:
 styles[n].font.name='Calibri';styles[n].font.size=Pt(sz);styles[n].font.color.rgb=RGBColor.from_string(NAVY);styles[n].paragraph_format.space_after=Pt(10)
styles['Caption'].font.size=Pt(10);styles['Caption'].font.color.rgb=RGBColor.from_string(GRAY)
styles['Normal'].paragraph_format.line_spacing=1.08
D.core_properties.title='VAS 5.0.0 — Operator and Administrator User Guide'
D.core_properties.subject='Verified interface workflows with synthetic demonstration data'
D.core_properties.author='VAS Documentation';D.core_properties.keywords='VAS, user guide, synthetic data, verified workflows'
sec=D.sections[0]
def setup(sec,land=False):
 sec.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
 sec.page_width=Inches(11.69 if land else 8.27);sec.page_height=Inches(8.27 if land else 11.69)
 sec.top_margin=Inches(.6);sec.bottom_margin=Inches(.6);sec.left_margin=Inches(.65);sec.right_margin=Inches(.65)
 sec.header_distance=Inches(.25);sec.footer_distance=Inches(.28)
setup(sec)
h=sec.header.paragraphs[0];h.text='VAS  /  USER GUIDE';h.style='Caption'
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('VAS 5.0.0  •  Verified 07 October 2026   |   ')
r=f.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
for r in f.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string(GRAY)
def p(text='',style=None):return D.add_paragraph(text,style)
def label(k,v):
 z=p();z.add_run(k+'  ').bold=True;z.add_run(v);return z
def heading(t,level=1,key=None):
 z=D.add_heading(t,level)
 if key:
  i=len(D.element.xpath('//w:bookmarkStart'))+1;b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(i));b.set(qn('w:name'),key);z._p.insert(0,b);e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(i));z._p.append(e)
 return z
def linkline(text,anchor):
 z=p();h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),anchor);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),TEAL);pr.append(co);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);z._p.append(h)
def newpage(land=False):
 s=D.add_section(WD_SECTION_START.NEW_PAGE);setup(s,land)
def bullets(items):
 for x in items:p(x,'List Bullet')
def table(headers,rows,widths=None):
 t=D.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
 for c,x in zip(t.rows[0].cells,headers):c.text=x
 trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for c,x in zip(cells,row):c.text=str(x)
 for row in t.rows:
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for z in c.paragraphs:
    z.paragraph_format.space_after=Pt(5)
    for r in z.runs:r.font.size=Pt(10)
 return t
p('OPERATOR + ADMINISTRATOR EDITION','Subtitle')
heading('VAS\nUser Guide',0)
p('Clear tasks. Real interface evidence. Explicit verification limits.','Subtitle')
p('Version 5.0.0  |  English  |  07 October 2026')
p('33 task walkthroughs • 71 discovered capabilities')
p('Genuine browser screenshots with editable numbered annotations')
z=p();z.alignment=WD_ALIGN_PARAGRAPH.CENTER
cover=z.add_run().add_picture(str(ROOT/'screenshots/72-recent-detections.png'),width=Inches(6.5))
cover._inline.docPr.set('descr','Genuine VAS Live Pipeline Feeds screen with synthetic silhouette identities; no real people.')
p('Prepared using an isolated local VAS instance and synthetic records. No real people, production data, passwords or tokens appear in this guide.')
label('Read first','“Tested” refers only to the stated demonstration. This guide does not certify recognition accuracy, GPU capacity, production security, backup recovery or model deployment.')
newpage();heading('Contents',1)
for n,title in sections.items():linkline(f'{n}. {title}  ................................  {PAGES.get("S"+str(n),"—")}',f'S{n}')
p('Use the task finder on the following pages to jump directly to an individual workflow. Orange numbered arrows match that workflow’s numbered instructions; numbering restarts for each task.')
p('Some figures show a control that was inspected but intentionally not activated. The verification label and limitations beside each task state exactly what was exercised.')
for j in range(0,len(WS),17):
 newpage();heading('Task finder'+(' — continued' if j else ''),1)
 for w in WS[j:j+17]:linkline(f'{w["id"]}  {w["title"]}  ·  p. {PAGES.get(w["id"],"—")}',w['id'])
newpage();heading('1. '+sections[1],1,'S1')
p('VAS supports review of face-related observations, identity histories, relationship analysis and operational administration. This guide separates what the application displayed from what was actually completed in a controlled demonstration.')
heading('Who should use this guide',2)
bullets(['Operators: review assigned camera feeds, identity evidence and searches.','Analysts: interpret relationships, patterns and risk evidence within granted access.','Administrators: manage accounts, locations, credentials, settings and service readiness.'])
heading('Before you begin',2)
bullets(['Obtain the approved VAS HTTPS address and an active account from your administrator. Server addresses are intentionally omitted here.','Confirm your assigned cameras and permitted tasks. A visible navigation item does not grant permission to every action.','Use approved demonstration images when learning. Do not enroll, export or alert on real people without organizational authorization.','Ask the administrator to confirm ingestion and required services are healthy. A running page does not establish that cameras, maps, training workers or the assistant are available.'])
heading('How to read the evidence',2)
table(['Status','Meaning'],[['Tested','The specified action and result were exercised in the isolated instance; read any qualifier.'],['Observed / untested','The control or screen was visible, but its consequential completion was not exercised.'],['Unavailable / blocked','A dependency, permission or usable dataset prevented completion in this run.'],['Not found','No dedicated capability was discovered in the reviewed VAS interface.']])
newpage();heading('Validation environment and boundaries',1)
label('Application','Runtime version 5.0.0; tutorial reports backend build 817b58e9460749bf53ef3a6fc12e024a6693d231 (development).')
label('Frontend','Current repository working tree: 52135d8-dirty. This is not a clean release-tag certification.')
label('Isolation','Separate database, cache, storage and internal container network; deployed API image reused with current checkout frontend. CPU execution. No production configuration was changed.')
label('Fixture','3 demonstration cameras, 14 synthetic identities and 210 seeded observations. Flat silhouette avatars are not face photographs. Similarities and detections are seeded examples, not model measurements.')
label('Accounts','Synthetic administrator and camera-scoped standard user. Password rotation and restricted access were tested. Observer and Analyzer roles were visible but not fully exercised.')
label('Missing dependencies','No training worker, offline basemap service or functioning assistant language model in this instance. No live RTSP source or sender ingestion test.')
label('Time','Browser display used Asia/Beirut. Dataset inspector explicitly uses UTC. Match the field’s stated timezone before interpreting ranges.')
label('Security','TLS and synthetic credentials were used for the private test. The test-only certificate setup is not a production installation recommendation.')
p('Screenshots preserve the actual interface, including genuine limitations. Password fields and one-time credentials were permanently raster-redacted before annotation. Original sanitized captures, focus crops, editable SVG overlays and JSON coordinates are included in the asset package.')

def intro(n):
 newpage();heading(f'{n}. {sections[n]}',1,f'S{n}')
 if n==2:
  p('Start at the approved VAS address. First login may require a password change. After rotation, verify the workspace and camera scope before continuing.')
  table(['Role shown in Create New User','Evidence and practical boundary'],[['ADMINISTRATOR - Full Access','Visible but disabled for assignment in this form. Existing synthetic administrator tested.'],['OBSERVER - Read-Only Access','Visible option; full permission matrix not exercised.'],['USER - Standard Access','Created with one assigned camera; restricted listing and administrator-route redirect verified.'],['ANALYZER - Analysis Access','Visible option; complete role workflow not exercised.']])
  p('Pipeline access and Chatbot access are separate controls. The system account is protected and is not an interactive user. Logout is available in the header; sign out when ending your session.')
 elif n==3:
  table(['Navigation','Purpose'],[['HOME','System and camera summaries'],['LIVE FEEDS','Latest stored detections and recent-detection panel'],['MANAGEMENT','People, users and pipeline administration'],['UNKNOWN FACES','Unrecognized identity records and evidence'],['SEARCH & INTELLIGENCE','Image search, history, relationship and security analysis'],['SYSTEM','Configuration, service operations and help']])
  p('Menus vary with authorization. “Live Pipeline Feeds” is a VAS detection display; these screenshots do not establish continuous video playback. VMS camera streaming and RTSP setup belong to a separate system.')
 elif n==5:p('Use the following tasks to review evidence, find identities, organize watchlists and configure an authorized alert rule. Always read the task’s verification boundary before reproducing a write action.')
 elif n==6:p('Administrator and analysis workflows follow. Do not infer administrator privileges from the /admin URL prefix alone: access depends on your account and camera scope. Settings and service changes require the applicable operational authorization.')
figcount=0
caption_overrides={
'08-operator':'Camera-scoped standard-user filter controls; scope was verified separately in the displayed results.',
'09-access-denied':'Standard-user result: an administrator URL redirects to the permitted Live Pipeline Feeds workspace.',
'38-membership-saved':'Observed display defect: the profile shows “Untitled watchlist”; membership was separately verified inside the named list.',
'42-search-quality':'Synthetic silhouette submitted for a negative quality test. No real face or positive recognition match was used.',
'28-map-unavailable':'Points or lines may render on an empty canvas; no geographic basemap was installed in the isolated instance.',
'30-ml-readiness':'Worker unavailable: training remains gated. This is readiness evidence, not a completed training run.',
'59-assistant-help':'Assistant interface/help entry in degraded mode; no language-model answer was validated.',
'44-credential-once':'One-time credential controls. The actual token is permanently redacted in the underlying screenshot.',
'62-retention-complete':'Execution-history filter used to locate the completed retention dry run; no destructive cleanup was initiated.'}
def figure(w,name):
 global figcount
 im=Image.open(ROOT/'annotated'/f'{name}.png');iw,ih=im.size
 land=iw/ih>1.55
 newpage(land);figcount+=1
 h=p(f'{w["id"]}  /  {w["title"]}','Heading 3');h.paragraph_format.keep_with_next=True
 marks=FIG[name]['marks'];nums=sorted(set(x['number'] for x in marks))
 part='Steps '+', '.join(map(str,nums)) if nums else 'Observed result / screen context'
 caption=caption_overrides.get(name,part+'. '+w['goal'])
 p(f'Figure {figcount:02d}. {caption}','Caption')
 maxw=10.39 if land else 6.97;maxh=5.5 if land else 8.5
 scale=min(maxw/iw,maxh/ih);z=p();z.alignment=WD_ALIGN_PARAGRAPH.CENTER
 pic=z.add_run().add_picture(str(ROOT/'annotated'/f'{name}.png'),width=Inches(iw*scale),height=Inches(ih*scale))
 pic._inline.docPr.set('descr',f'Genuine sanitized VAS screenshot. {w["title"]}. {part}. {caption}')
 pic._inline.docPr.set('title',f'{w["id"]} figure {figcount:02d} — {name}')
 z.paragraph_format.space_after=Pt(1)
 note=p('Orange numbers refer to this task’s written steps. Synthetic demonstration data.','Caption');note.paragraph_format.space_after=Pt(0)
def workflow(w):
 newpage();heading(w['id']+'  '+w['title'],1,w['id'])
 label('Accomplish',w['goal']);label('Permissions / prerequisites',w['role']);label('Start',w['start']);label('Verification',w['status'])
 for i,s in enumerate(w['steps'],1):
  z=p();z.add_run(f'{i}.  ').bold=True;z.add_run(s)
 label('Expected result / verify',w['expected']);label('Errors / recovery',w['recovery'])
 if w['limit']:label('Boundary',w['limit'])
 for name in w['figures']:figure(w,name)
for n in [2,3]:
 intro(n)
 for w in WS:
  if w['section']==n:workflow(w)
newpage();heading('4. '+sections[4],1,'S4')
p('First-session route: sign in → confirm camera scope → inspect an identity → review its recorded journey → verify the evidence. This quick start links to the exact annotated tasks; it does not create new cameras or train a model.')
for text,wid in [('Sign in and complete required password rotation','W01'),('Review Home and the camera selector','W04'),('Open Unknown Faces and apply a camera filter','W06'),('Open VIEW & ALERT and inspect the recorded movement journey','W07'),('Review image search prerequisites and quality checks','W10'),('Create a watchlist only when authorized','W12')]:linkline(text+'  →  '+wid,wid)
heading('Successful first session',2)
p('You can access the intended camera scope, open an identity record and follow timestamped appearances. Confirm the displayed camera and time at each step. A silhouette, empty result, stale feed or insufficient-baseline message is evidence to investigate, not evidence of a successful recognition or analysis.')
heading('Administrator handoff',2)
p('Provide users with the approved address, initial account, assigned cameras and support contact. Verify the sender-to-VAS ingestion path, storage access and optional dependencies separately before operational use. Use W16 for camera-scoped account creation and W25–W29 for analytics readiness.')
for n in [5,6]:
 intro(n)
 for w in WS:
  if w['section']==n:workflow(w)
newpage();heading('7. '+sections[7],1,'S7')
issues=[
('Login blocked or repeated sign-in','Confirm the approved HTTPS address, account status and required rotation. Ask the administrator to inspect origin/session policy. Never publish passwords in screenshots.','W01–W02'),
('A page redirects or cameras are missing','Check role and Pipeline access. The tested standard user was redirected away from user administration and saw only assigned camera evidence.','W16'),
('A feed looks empty or stale','Confirm ingestion, latest observation time and assigned camera. VAS stored detection display does not prove that the VMS browser video transport is healthy.','W05'),
('Image is rejected or search returns no faces','Use an approved clear face image and Quality Check. Synthetic silhouette input correctly produced “No face was detected in the image.” Positive matching was not tested.','W10'),
('Dialog action labels have poor contrast','Some watchlist and alert submit controls rendered dark text on a dark background. The enlarged figures point to their actual positions; this UI defect needs a separate application fix.','W13–W14'),
('A saved watchlist looks unnamed','The profile chip showed “Untitled watchlist” in this run. Verify the identity in the named watchlist’s View panel before concluding the write failed.','W13'),
('Alert rule exists but no alert arrives','Inspect Health, rule state, similarity and delivery settings. Rule save/pause was tested; triggering, sound, email and SMS were not. The fixture snapshot health check reported a problem.','W14'),
('Anomaly detection has insufficient baseline','The comparison interval needs enough earlier observations. The 90-day example had no earlier baseline; gather valid history or use a justified shorter interval. Do not weaken validation merely to obtain a score.','W23'),
('Threat score seems decisive','Read Requested mode, Executed mode, risk factors and scoring provenance. The demonstrated rules score is not a probability that someone is a threat.','W24'),
('Map is blank','Confirm offline basemap installation and service availability. Coordinates or trajectory lines alone do not establish a functioning map or security-zone geometry.','W26'),
('Prepare & train is disabled','Restore worker availability and satisfy source/dataset readiness. Keep minimum rows, split, null and leakage checks. Do not treat a ready page as a deployed model.','W27'),
('Dataset preview does not load','The isolated inspector had no camera choices and no preview output. Report the affected service, range and screen state; no notebook execution was validated here.','W28'),
('Similarity training asks for more samples','Collect approved and rejected feedback for sufficient distinct pairs. The displayed baseline was 0/50 samples, 0/5 approved, 0/5 rejected and 0/10 pairs.','W29'),
('Assistant says degraded','Ask the administrator to check its model service and permissions. No assistant answer was tested in this instance.','W30'),
('Retention change appears delayed','Read the apply timing. The tested retention setting applied at the next scheduled cleanup; use the dry test to preview eligibility without deletion.','W18–W19')]
for i in range(0,len(issues),4):
 if i:newpage();heading('Troubleshooting — continued',1)
 for title,body,ref in issues[i:i+4]:heading(title,2);p(body);p('Related task: '+ref,'Caption')
newpage();heading('8. '+sections[8],1,'S8')
p('The companion Capability-coverage.csv contains the full feature inventory: purpose, location, permission, verification status, task reference, screenshot evidence and remaining limitations. The following register includes every discovered capability from that inventory.')
heading('What was established',2)
p('Verified demonstrations include authentication and rotation, camera-scoped access, location and name updates, filtering and identity journeys, watchlist creation/membership, alert save/health/pause, CSV export, credential issuance, settings persistence, a completed retention dry run, relationship/pattern generation and analysis-readiness behavior.')
heading('What remains unverified',2)
bullets(['Positive face recognition, real-photo enrollment, merge and promotion completion; live ingestion and camera throughput.','Triggered notifications and external delivery; destructive cleanup, record deletion, credential revocation and recovery from backups.','End-to-end ML training, notebook execution, hyperparameter tuning, drift monitoring and model connection to consuming services.','Working basemaps/security-zone geometry and assistant answers; complete Observer/Analyzer permission coverage.'])
heading('Features not discovered as dedicated VAS screens',2)
p('Continuous video playback or recording archive, vehicle/license-plate operations, multi-site orchestration and a general report builder were not found. Search-history CSV export and timestamped face appearances do exist. Do not substitute those for a video archive or validated reporting subsystem.')
heading('Interpretation',2)
p('The inventory is an evidence record for this application state, not a claim that every listed capability is production ready. Re-run the affected tasks after releases, configuration changes or dependency restoration.')
for i in range(0,len(COV),4):
 newpage();heading(f'Capability register · {i+1}–{min(i+4,len(COV))}',1)
 for j,c in enumerate(COV[i:i+4],i+1):
  heading(f'C{j:02d}  {c["feature"]}',2)
  p(c['purpose']);label('Location / permission',c['location']+'  |  '+c['permissions']);label('Status / guide',c['verification_status']+'  |  '+c['guide_section'])
  if c['limitations']:label('Limit',c['limitations'])
  if c['evidence']:p('Capture evidence: '+str(c['evidence']),'Caption')
newpage();heading('Maintaining this guide',1)
p('Edit source/workflows.json and source/coverage.json when behavior changes. Retest against an authorized isolated instance; replace screenshots with sanitized genuine captures. JSON annotation coordinates and SVG groups are editable. Re-render the Word file to PDF and inspect every page after edits.')
p('Keep status qualifiers. A screenshot of a disabled button, an empty table or an unavailable dependency is useful evidence, but it is not a completed workflow.')
p('Delivery contents: editable DOCX, matching PDF, sanitized screenshot folder, annotated PNGs, self-contained editable SVGs, annotation JSON, capability CSV, canonical content JSON, build scripts and the verification report.')
D.save(ROOT/'VAS-User-Guide.docx')
print('Built Word guide:',len(WS),'tasks,',figcount,'figures,',len(COV),'capabilities')
