"""The MCP tool catalogue: narrow tools with validated input and output.

Every tool:
  * has a Pydantic input model (unknown fields refused) and an output model;
  * returns a structured envelope: ``{status, tool, version, data, error_code,
    error}`` - never raises into the caller;
  * runs under a timeout;
  * emits ``fr_agent_stage_duration_seconds{stage="mcp_call", component=<tool>}``;
  * writes one audit line (tool name, outcome, duration, request id, user id -
    never arguments, never rows, never SQL text);
  * leaks no secret: connection strings, keys and credentials never enter a
    result.

The database tools sit BEHIND the existing controls: ``sql.validate`` and
``security.*`` are the AST guard (``sql_agent/security/sql_guard.py``),
``database.execute_readonly`` is ``DatabaseManager.execute_query`` with the
caller's camera scope and the read-only role. No tool accepts arbitrary SQL
for execution without that gate, and no ``execute_any_sql`` exists.
"""
from __future__ import annotations

import csv
import io
import logging
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor, TimeoutError as FutureTimeout
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, List, Optional

from pydantic import BaseModel, ConfigDict, Field, ValidationError

logger = logging.getLogger(__name__)

VERSION = "1.0"
ERROR_CODES = ("INVALID_ARGUMENTS", "INVALID_RESULT", "TIMEOUT",
               "DEPENDENCY_UNAVAILABLE", "PERMISSION_DENIED", "NOT_SUPPORTED", "INTERNAL")


class ToolError(Exception):
    """A structured failure. ``detail`` carries machine facts a caller may
    act on (for a refused statement: the guard's own code and message)."""

    def __init__(self, code: str, message: str, detail=None):
        assert code in ERROR_CODES, code
        super().__init__(message)
        self.code = code
        self.message = message
        self.detail = dict(detail or {})


@dataclass
class ToolContext:
    """Who is calling. Scope comes from the caller's session, never from the model."""
    user_id: Optional[int] = None
    role: Optional[str] = None
    request_id: str = "-"
    pipeline_scope: Optional[frozenset] = None     # None = administrator


class _In(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=False)


class _Out(BaseModel):
    model_config = ConfigDict(extra="forbid")


# ---------------------------------------------------------------- schema I/O
class SearchSchemaIn(_In):
    query: str = Field(min_length=1, max_length=400)
    top_k: int = Field(default=5, ge=1, le=20)


class SchemaHit(_Out):
    table: str
    column: Optional[str] = None
    description: str = ""
    score: float


class SearchSchemaOut(_Out):
    hits: List[SchemaHit]


class DescribeTableIn(_In):
    table: str = Field(min_length=1, max_length=64)


class ColumnInfo(_Out):
    name: str
    type: str
    nullable: bool
    description: str = ""


class DescribeTableOut(_Out):
    table: str
    description: str = ""
    columns: List[ColumnInfo]
    primary_keys: List[str] = Field(default_factory=list)
    foreign_keys: List[Dict[str, str]] = Field(default_factory=list)


class Empty(_In):
    pass


class RelationshipsOut(_Out):
    relationships: List[str]


class AllowedTablesOut(_Out):
    tables: List[str]


# ---------------------------------------------------------------- vanna I/O
class RetrieveExamplesIn(_In):
    question: str = Field(min_length=1, max_length=1000)
    top_k: int = Field(default=5, ge=1, le=20)


class SqlExample(_Out):
    question: str
    sql: str
    purpose: str = ""
    similarity: float
    source: str


class RetrieveExamplesOut(_Out):
    examples: List[SqlExample]


class RetrieveDdlIn(_In):
    tables: List[str] = Field(default_factory=list, max_length=20)


class RetrieveDdlOut(_Out):
    ddl: str


class BusinessContextIn(_In):
    question: str = Field(min_length=1, max_length=1000)
    top_k: int = Field(default=6, ge=1, le=30)


class BusinessContextOut(_Out):
    notes: List[str]


# ------------------------------------------------------------------ sql I/O
class SqlIn(_In):
    sql: str = Field(min_length=1, max_length=20000)


class ValidateOut(_Out):
    allowed: bool
    code: str
    reason: str = ""
    statement_type: str = ""
    tables: List[str] = Field(default_factory=list)
    canonical_sql: str = ""


class ExplainOut(_Out):
    statement_type: str
    tables: List[str]
    filters: List[str]
    aggregates: List[str]
    group_by: List[str]
    order_by: List[str]
    limit: Optional[int] = None
    summary: str


class GenerateIn(_In):
    question: str = Field(min_length=1, max_length=1000)


class GenerateOut(_Out):
    sql: str
    purpose: str = ""


class ApplyLimitsIn(_In):
    sql: str = Field(min_length=1, max_length=20000)
    max_rows: int = Field(default=500, ge=1, le=5000)


class ApplyLimitsOut(_Out):
    sql: str
    max_rows: int


class ReadonlyOut(_Out):
    read_only: bool
    statement_type: str
    reason: str = ""


class AllowedCheckOut(_Out):
    allowed: bool
    referenced: List[str]
    unauthorized: List[str]


# ------------------------------------------------------------- database I/O
class ExecuteIn(_In):
    sql: str = Field(min_length=1, max_length=20000)
    max_rows: int = Field(default=500, ge=1, le=5000)


class ExecuteOut(_Out):
    columns: List[str]
    rows: List[Dict[str, Any]]
    row_count: int
    truncated: bool = False
    duration_ms: int


class ExplainQueryOut(_Out):
    plan: Any


class CancelIn(_In):
    request_id: str = Field(min_length=1, max_length=64)


class CancelOut(_Out):
    cancelled: bool


# ------------------------------------------------------------ analytics I/O
class AggregateIn(_In):
    rows: List[Dict[str, Any]] = Field(max_length=5000)
    group_by: Optional[str] = None
    metric: Optional[str] = None
    op: str = Field(default="count", pattern="^(count|sum|avg|min|max)$")


class AggregateOut(_Out):
    groups: List[Dict[str, Any]]


class CompareIn(_In):
    rows: List[Dict[str, Any]] = Field(max_length=5000)
    key: str
    left: str
    right: str
    value: str


class CompareOut(_Out):
    left: Dict[str, Any]
    right: Dict[str, Any]
    difference: float
    percentage_difference: Optional[float] = None


class TrendIn(_In):
    rows: List[Dict[str, Any]] = Field(max_length=5000)
    x: str
    y: str


class TrendOut(_Out):
    points: List[Dict[str, Any]]
    first: Optional[float] = None
    last: Optional[float] = None
    change: Optional[float] = None
    direction: str


class PctChangeIn(_In):
    before: float
    after: float


class PctChangeOut(_Out):
    before: float
    after: float
    change: float
    percentage_change: Optional[float] = None


# ------------------------------------------------------------ reporting I/O
class ReportIn(_In):
    title: str = Field(min_length=1, max_length=200)
    content: str = Field(default="", max_length=200000)
    rows: List[Dict[str, Any]] = Field(default_factory=list, max_length=5000)
    columns: List[str] = Field(default_factory=list, max_length=100)


class ReportOut(_Out):
    format: str
    bytes_base64: str
    filename: str


# --------------------------------------------------------------- system I/O
class HealthOut(_Out):
    database: bool
    knowledge_base_examples: int
    providers: List[str]
    offline_mode: bool


class CapabilitiesOut(_Out):
    tools: List[str]
    version: str
    offline_mode: bool
    vector_store: str


# ================================================================= the toolset
@dataclass
class ToolSpec:
    name: str
    description: str
    input_model: type
    output_model: type
    handler: Callable[..., Any]
    timeout_seconds: float = 30.0
    reads_data: bool = False


_WRITE_STATEMENTS = ("INSERT", "UPDATE", "DELETE", "DROP", "ALTER", "TRUNCATE", "CREATE",
                     "MERGE", "REPLACE", "GRANT", "REVOKE", "EXECUTE", "CALL", "COPY")


def _tokens(text: str) -> set:
    return {t for t in re.split(r"[^a-z0-9]+", (text or "").lower()) if len(t) > 2}


class MCPToolset:
    """The catalogue, bound to one caller's database manager and knowledge base."""

    def __init__(self, db=None, kb=None, *, generator: Optional[Callable[[str], Dict[str, str]]] = None,
                 canceller: Optional[Callable[[str], bool]] = None, config=None,
                 providers: Optional[List[str]] = None):
        self.db = db
        self.kb = kb
        self.generator = generator
        self.canceller = canceller
        self.config = config
        self.providers = list(providers or [])
        self._pool = ThreadPoolExecutor(max_workers=4, thread_name_prefix="mcp-tool")
        self._specs: Dict[str, ToolSpec] = {}
        self._register_all()

    # ------------------------------------------------------------- registry
    def _register(self, name, description, input_model, output_model, handler,
                  timeout_seconds=30.0, reads_data=False):
        self._specs[name] = ToolSpec(name, description, input_model, output_model, handler,
                                     timeout_seconds, reads_data)

    def specs(self) -> List[ToolSpec]:
        return [self._specs[n] for n in sorted(self._specs)]

    def names(self) -> List[str]:
        return sorted(self._specs)

    def json_schemas(self) -> List[Dict[str, Any]]:
        """OpenAI-style function specs, for any orchestrator."""
        return [{"name": s.name, "description": s.description,
                 "parameters": s.input_model.model_json_schema()} for s in self.specs()]

    # ----------------------------------------------------------------- call
    def call(self, name: str, arguments: Optional[Dict[str, Any]] = None,
             ctx: Optional[ToolContext] = None) -> Dict[str, Any]:
        """One tool call, traced as an Opik span in development."""
        runner = self._tracked_runner(name)
        return runner(name, arguments, ctx)

    def _tracked_runner(self, name: str):
        cache = self.__dict__.setdefault("_runners", {})
        runner = cache.get(name)
        if runner is None:
            runner = self._call
            try:
                from .. import tracing
                decorator = tracing.tool_tracker(self.config, name) if self.config is not None else None
                if decorator is not None:
                    runner = decorator(self._call)
            except Exception:
                runner = self._call
            cache[name] = runner
        return runner

    def _call(self, name: str, arguments: Optional[Dict[str, Any]] = None,
              ctx: Optional[ToolContext] = None) -> Dict[str, Any]:
        ctx = ctx or ToolContext()
        started = time.monotonic()
        spec = self._specs.get(name)
        if spec is None:
            return self._finish(name, ctx, started, error=ToolError("NOT_SUPPORTED", f"unknown tool {name!r}"))
        try:
            arguments = spec.input_model.model_validate(arguments or {})
        except ValidationError as e:
            return self._finish(name, ctx, started, error=ToolError(
                "INVALID_ARGUMENTS", "; ".join(f"{'.'.join(str(p) for p in err['loc'])}: {err['msg']}"
                                               for err in e.errors()[:5])))
        future = self._pool.submit(spec.handler, arguments, ctx)
        try:
            raw = future.result(timeout=spec.timeout_seconds)
        except FutureTimeout:
            future.cancel()
            return self._finish(name, ctx, started, error=ToolError(
                "TIMEOUT", f"{name} exceeded {spec.timeout_seconds:.0f}s"))
        except ToolError as e:
            return self._finish(name, ctx, started, error=e)
        except Exception as e:  # never leak an exception (or a DSN inside it)
            logger.warning("[MCP] %s failed: %s", name, type(e).__name__)
            return self._finish(name, ctx, started, error=ToolError(
                "DEPENDENCY_UNAVAILABLE", f"{name} failed: {type(e).__name__}"))
        try:
            data = spec.output_model.model_validate(raw).model_dump()
        except ValidationError as e:
            logger.warning("[MCP] %s produced an invalid result: %s", name, e.errors()[:1])
            return self._finish(name, ctx, started, error=ToolError(
                "INVALID_RESULT", f"{name} returned data outside its contract"))
        return self._finish(name, ctx, started, data=data)

    def _finish(self, name, ctx, started, *, data=None, error: Optional[ToolError] = None):
        seconds = time.monotonic() - started
        outcome = "ok" if error is None else ("timeout" if error.code == "TIMEOUT"
                                              else "rejected" if error.code in ("INVALID_ARGUMENTS", "PERMISSION_DENIED")
                                              else "error")
        try:
            from .. import observability
            observability.observe_stage_latency("mcp_call", seconds, outcome, name)
        except Exception:
            pass
        try:
            from backend.auth.auth_security import audit
            audit("mcp_tool_call", result="success" if error is None else "failure",
                  request_id=str(ctx.request_id or "-"), user_id=ctx.user_id,
                  failure_code=error.code if error else None,
                  duration_ms=int(seconds * 1000), tool=name, role=ctx.role)
        except Exception:
            pass
        envelope = {"status": "ok" if error is None else "error", "tool": name, "version": VERSION,
                    "data": data or {}, "error_code": error.code if error else None,
                    "error": error.message if error else None,
                    "detail": dict(error.detail) if error is not None and error.detail else {}}
        return envelope

    # ------------------------------------------------------- shared helpers
    def _schema(self) -> Dict[str, Any]:
        if self.db is None:
            raise ToolError("DEPENDENCY_UNAVAILABLE", "no database manager bound")
        return getattr(self.db, "KNOWN_SCHEMA", None) or {"tables": {}, "relationships": []}

    def _policy(self, ctx: ToolContext):
        from ..security import sql_guard
        if self.db is None or getattr(self.db, "sql_policy", None) is None:
            raise ToolError("DEPENDENCY_UNAVAILABLE", "no SQL policy bound")
        import dataclasses
        policy = self.db.sql_policy
        if ctx.pipeline_scope is not None:
            policy = dataclasses.replace(policy, pipeline_scope=frozenset(str(p) for p in ctx.pipeline_scope))
        return sql_guard, policy

    def _verdict(self, sql: str, ctx: ToolContext):
        sql_guard, policy = self._policy(ctx)
        return sql_guard.validate_sql(sql, policy)

    # --------------------------------------------------------------- schema
    def _search_schema(self, args: SearchSchemaIn, ctx):
        words = _tokens(args.query)
        hits = []
        for table, info in self._schema().get("tables", {}).items():
            base = _tokens(table + " " + info.get("description", ""))
            score = len(words & base) / max(1, len(words))
            if score:
                hits.append({"table": table, "column": None, "description": info.get("description", ""), "score": round(score, 3)})
            for col in info.get("columns", []):
                cwords = _tokens(col.get("column_name", "") + " " + col.get("description", ""))
                cscore = len(words & cwords) / max(1, len(words))
                if cscore:
                    hits.append({"table": table, "column": col.get("column_name"),
                                 "description": col.get("description", ""), "score": round(cscore, 3)})
        hits.sort(key=lambda h: (-h["score"], h["table"], h["column"] or ""))
        return {"hits": hits[:args.top_k]}

    def _describe_table(self, args: DescribeTableIn, ctx):
        tables = self._schema().get("tables", {})
        info = tables.get(args.table.lower())
        if info is None:
            raise ToolError("PERMISSION_DENIED", f"table {args.table!r} is not exposed to the agent")
        return {"table": args.table.lower(), "description": info.get("description", ""),
                "columns": [{"name": c.get("column_name"), "type": c.get("data_type", ""),
                             "nullable": str(c.get("is_nullable", "YES")).upper() != "NO",
                             "description": c.get("description", "")} for c in info.get("columns", [])],
                "primary_keys": list(info.get("primary_keys", [])),
                "foreign_keys": [{k: str(v) for k, v in fk.items()} for fk in info.get("foreign_keys", [])]}

    def _relationships(self, args, ctx):
        return {"relationships": list(self._schema().get("relationships", []))}

    def _allowed_tables(self, args, ctx):
        _, policy = self._policy(ctx)
        return {"tables": sorted(policy.allowed_tables)}

    # ---------------------------------------------------------------- vanna
    def _retrieve_examples(self, args: RetrieveExamplesIn, ctx):
        if self.kb is None:
            raise ToolError("DEPENDENCY_UNAVAILABLE", "no knowledge base bound")
        examples = self.kb.search_similar(query=args.question, top_k=args.top_k, user_id=ctx.user_id) or []
        return {"examples": [{"question": e.get("question", ""), "sql": e.get("sql", ""),
                              "purpose": e.get("purpose", "") or "", "similarity": float(e.get("similarity") or 0.0),
                              "source": str(e.get("source") or "unknown")} for e in examples]}

    def _retrieve_ddl(self, args: RetrieveDdlIn, ctx):
        tables = self._schema().get("tables", {})
        wanted = [t.lower() for t in args.tables] or sorted(tables)
        parts = []
        for name in wanted:
            info = tables.get(name)
            if info is None:
                continue
            cols = ",\n".join(f"    {c.get('column_name')} {c.get('data_type', '')}"
                              + ("" if str(c.get("is_nullable", "YES")).upper() != "NO" else " NOT NULL")
                              + (f"  -- {c.get('description')}" if c.get("description") else "")
                              for c in info.get("columns", []))
            fks = "".join(f",\n    FOREIGN KEY ({fk.get('column_name')}) REFERENCES "
                          f"{fk.get('foreign_table_name')}({fk.get('foreign_column_name')})"
                          for fk in info.get("foreign_keys", []))
            pk = ",\n    PRIMARY KEY (" + ", ".join(info.get("primary_keys", [])) + ")" if info.get("primary_keys") else ""
            parts.append(f"-- {info.get('description', '')}\nCREATE TABLE {name} (\n{cols}{pk}{fks}\n);")
        return {"ddl": "\n\n".join(parts)}

    def _business_context(self, args: BusinessContextIn, ctx):
        words = _tokens(args.question)
        notes = list(self._schema().get("relationships", []))
        ranked = sorted(notes, key=lambda n: -len(words & _tokens(n)))
        return {"notes": ranked[:args.top_k]}

    # ------------------------------------------------------------------ sql
    def _validate(self, args: SqlIn, ctx):
        v = self._verdict(args.sql, ctx)
        return {"allowed": bool(v.allowed), "code": v.code, "reason": v.reason or "",
                "statement_type": v.statement_type or "", "tables": list(v.tables or []),
                "canonical_sql": v.canonical or v.sql or ""}

    def _explain(self, args: SqlIn, ctx):
        try:
            import sqlglot
            from sqlglot import expressions as exp
            tree = sqlglot.parse_one(args.sql, read="postgres")
        except Exception as e:
            raise ToolError("INVALID_ARGUMENTS", f"SQL does not parse: {type(e).__name__}")
        ctes = {c.alias_or_name.lower() for c in tree.find_all(exp.CTE)}
        tables = sorted({t.name.lower() for t in tree.find_all(exp.Table) if t.name and t.name.lower() not in ctes})
        filters = [w.this.sql(dialect="postgres") for w in tree.find_all(exp.Where)]
        aggregates = sorted({a.sql(dialect="postgres") for a in tree.find_all(exp.AggFunc)})[:20]
        group_by = [g.sql(dialect="postgres") for grp in tree.find_all(exp.Group) for g in grp.expressions]
        order_by = [o.sql(dialect="postgres") for ordr in tree.find_all(exp.Order) for o in ordr.expressions]
        limit_node = tree.args.get("limit") if isinstance(tree, exp.Select) else None
        limit = None
        try:
            limit = int(limit_node.expression.this) if limit_node is not None else None
        except Exception:
            limit = None
        kind = type(tree).__name__.upper()
        summary = (f"{kind} over {', '.join(tables) or 'no tables'}"
                   + (f", filtered by {len(filters)} condition(s)" if filters else "")
                   + (f", aggregated with {', '.join(aggregates[:3])}" if aggregates else "")
                   + (f", grouped by {', '.join(group_by[:3])}" if group_by else "")
                   + (f", limited to {limit} rows" if limit else ""))
        return {"statement_type": kind, "tables": tables, "filters": filters[:10], "aggregates": aggregates,
                "group_by": group_by[:10], "order_by": order_by[:10], "limit": limit, "summary": summary}

    def _generate(self, args: GenerateIn, ctx):
        if self.generator is None:
            raise ToolError("NOT_SUPPORTED", "sql.generate is served by the orchestrator's generation "
                                             "node in this deployment; call the agent, not the tool")
        out = self.generator(args.question) or {}
        if not out.get("sql"):
            raise ToolError("DEPENDENCY_UNAVAILABLE", "the generator produced no SQL")
        return {"sql": out["sql"], "purpose": out.get("purpose", "")}

    # ------------------------------------------------------------- security
    def _authorize(self, args: SqlIn, ctx):
        return self._validate(args, ctx)

    def _enforce_readonly(self, args: SqlIn, ctx):
        upper = " ".join(args.sql.split()).upper().strip().rstrip(";")
        if ";" in upper:
            return {"read_only": False, "statement_type": "MULTIPLE", "reason": "stacked statements are refused"}
        try:
            import sqlglot
            from sqlglot import expressions as exp
            tree = sqlglot.parse_one(args.sql, read="postgres")
        except Exception as e:
            return {"read_only": False, "statement_type": "UNPARSEABLE", "reason": f"{type(e).__name__}"}
        kind = type(tree).__name__.upper()
        if not isinstance(tree, (exp.Select, exp.Union)) or any(
                isinstance(n, (exp.Insert, exp.Update, exp.Delete, exp.Drop, exp.Alter, exp.Create,
                               exp.Merge, exp.Command, exp.TruncateTable)) for n in tree.walk()):
            return {"read_only": False, "statement_type": kind, "reason": "only SELECT is permitted"}
        for word in _WRITE_STATEMENTS:
            if re.search(rf"\b{word}\b", upper) and word not in ("CREATE",) and kind not in ("SELECT", "UNION"):
                return {"read_only": False, "statement_type": kind, "reason": f"{word} is refused"}
        return {"read_only": True, "statement_type": kind, "reason": ""}

    def _check_allowed_tables(self, args: SqlIn, ctx):
        v = self._verdict(args.sql, ctx)
        _, policy = self._policy(ctx)
        referenced = sorted(set(v.tables or []))
        unauthorized = sorted(t for t in referenced if t.split(".")[-1] not in policy.allowed_tables)
        return {"allowed": not unauthorized and v.code != "TABLE_NOT_ALLOWED",
                "referenced": referenced, "unauthorized": unauthorized}

    def _apply_limits(self, args: ApplyLimitsIn, ctx):
        import dataclasses
        sql_guard, policy = self._policy(ctx)
        policy = dataclasses.replace(policy, max_rows=min(int(args.max_rows), int(policy.max_rows)))
        v = sql_guard.validate_sql(args.sql, policy)
        if not v.allowed:
            raise ToolError("PERMISSION_DENIED", f"{v.code}: {v.reason}")
        return {"sql": v.sql or args.sql, "max_rows": policy.max_rows}

    # ------------------------------------------------------------- database
    def _execute_readonly(self, args: ExecuteIn, ctx):
        if self.db is None:
            raise ToolError("DEPENDENCY_UNAVAILABLE", "no database manager bound")
        if getattr(self.db, "sql_policy", None) is not None:
            verdict = self._verdict(args.sql, ctx)
            if not verdict.allowed:
                raise ToolError("PERMISSION_DENIED", f"{verdict.code}: {verdict.reason}",
                                detail={"error_code": verdict.code, "error": f"Security: {verdict.reason}"})
        # DatabaseManager.execute_query runs the AST guard and the scope again
        # as its first statement, so a manager without a bound policy (a test
        # double) is still never handed anything unguarded by the real code.
        started = time.monotonic()
        result = self.db.execute_query(args.sql) or {}
        if not result.get("success"):
            first_line = str(result.get("error") or "execution failed").splitlines()[0]
            # The manager's own denial (its guard runs first) keeps its code:
            # the loop's enforcement distinguishes a parse error from DELETE.
            if result.get("error_code"):
                raise ToolError("PERMISSION_DENIED", f"{result['error_code']}: {first_line[:200]}",
                                detail={"error_code": result["error_code"], "error": first_line[:300]})
            raise ToolError("DEPENDENCY_UNAVAILABLE", first_line[:200], detail={"error": first_line[:300]})
        rows = list(result.get("rows") or [])[:args.max_rows]
        return {"columns": list(result.get("columns") or (list(rows[0].keys()) if rows else [])),
                "rows": rows, "row_count": int(result.get("row_count") or len(rows)),
                "truncated": bool(result.get("truncated")) or len(result.get("rows") or []) > len(rows),
                "duration_ms": int((time.monotonic() - started) * 1000)}

    def _explain_query(self, args: SqlIn, ctx):
        _, policy = self._policy(ctx)
        if not getattr(policy, "allow_explain", False):
            raise ToolError("PERMISSION_DENIED", "EXPLAIN is disabled by the SQL policy")
        verdict = self._verdict(args.sql, ctx)
        if not verdict.allowed:
            raise ToolError("PERMISSION_DENIED", f"{verdict.code}: {verdict.reason}")
        result = self.db.execute_query("EXPLAIN (FORMAT JSON) " + (verdict.sql or args.sql)) or {}
        if not result.get("success"):
            raise ToolError("DEPENDENCY_UNAVAILABLE", str(result.get("error") or "explain failed")[:200])
        return {"plan": result.get("rows")}

    def _cancel(self, args: CancelIn, ctx):
        if self.canceller is None:
            raise ToolError("NOT_SUPPORTED", "cancellation is served by the API layer")
        return {"cancelled": bool(self.canceller(args.request_id))}

    # ------------------------------------------------------------ analytics
    @staticmethod
    def _num(value) -> Optional[float]:
        try:
            return float(value)
        except (TypeError, ValueError):
            return None

    def _aggregate(self, args: AggregateIn, ctx):
        groups: Dict[Any, List[float]] = {}
        for row in args.rows:
            key = row.get(args.group_by) if args.group_by else "all"
            value = self._num(row.get(args.metric)) if args.metric else 1.0
            groups.setdefault(key, []).append(value if value is not None else 0.0)
        out = []
        for key, values in groups.items():
            if args.op == "count":
                figure = float(len(values))
            elif args.op == "sum":
                figure = float(sum(values))
            elif args.op == "avg":
                figure = float(sum(values) / len(values)) if values else 0.0
            elif args.op == "min":
                figure = float(min(values)) if values else 0.0
            else:
                figure = float(max(values)) if values else 0.0
            out.append({"group": key, "op": args.op, "value": round(figure, 4)})
        out.sort(key=lambda g: -g["value"])
        return {"groups": out}

    def _compare(self, args: CompareIn, ctx):
        left = next((r for r in args.rows if str(r.get(args.key)) == args.left), None)
        right = next((r for r in args.rows if str(r.get(args.key)) == args.right), None)
        if left is None or right is None:
            raise ToolError("INVALID_ARGUMENTS", "left or right key not found in rows")
        lv, rv = self._num(left.get(args.value)) or 0.0, self._num(right.get(args.value)) or 0.0
        pct = round(100.0 * (lv - rv) / rv, 2) if rv else None
        return {"left": {args.key: args.left, args.value: lv}, "right": {args.key: args.right, args.value: rv},
                "difference": round(lv - rv, 4), "percentage_difference": pct}

    def _trend(self, args: TrendIn, ctx):
        points = sorted(({"x": r.get(args.x), "y": self._num(r.get(args.y))} for r in args.rows
                         if r.get(args.x) is not None), key=lambda p: str(p["x"]))
        ys = [p["y"] for p in points if p["y"] is not None]
        first, last = (ys[0], ys[-1]) if ys else (None, None)
        change = round(last - first, 4) if ys else None
        direction = "flat" if not ys or change == 0 else ("up" if change > 0 else "down")
        return {"points": points, "first": first, "last": last, "change": change, "direction": direction}

    def _pct_change(self, args: PctChangeIn, ctx):
        change = round(args.after - args.before, 4)
        pct = round(100.0 * change / args.before, 2) if args.before else None
        return {"before": args.before, "after": args.after, "change": change, "percentage_change": pct}

    # ------------------------------------------------------------ reporting
    def _csv(self, args: ReportIn, ctx):
        import base64
        buf = io.StringIO()
        columns = args.columns or (list(args.rows[0].keys()) if args.rows else [])
        writer = csv.DictWriter(buf, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in args.rows:
            writer.writerow({c: row.get(c) for c in columns})
        data = buf.getvalue().encode("utf-8")
        return {"format": "csv", "bytes_base64": base64.b64encode(data).decode("ascii"),
                "filename": re.sub(r"[^A-Za-z0-9_-]+", "_", args.title)[:60] + ".csv"}

    def _pdf(self, args: ReportIn, ctx):
        import base64
        from datetime import datetime
        try:
            from ..services.export_builders import build_pdf_bytes
        except Exception:
            raise ToolError("DEPENDENCY_UNAVAILABLE", "PDF builder not available")
        content = args.content or "\n".join(", ".join(f"{k}: {v}" for k, v in r.items()) for r in args.rows)
        data = build_pdf_bytes(args.title, content, datetime.utcnow().strftime("%Y-%m-%d"), "FACE_DETECTOR")
        return {"format": "pdf", "bytes_base64": base64.b64encode(data).decode("ascii"),
                "filename": re.sub(r"[^A-Za-z0-9_-]+", "_", args.title)[:60] + ".pdf"}

    def _excel(self, args: ReportIn, ctx):
        import base64
        try:
            import openpyxl
        except ImportError:
            raise ToolError("DEPENDENCY_UNAVAILABLE", "openpyxl is not installed")
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = args.title[:31] or "report"
        columns = args.columns or (list(args.rows[0].keys()) if args.rows else [])
        sheet.append(columns)
        for row in args.rows:
            sheet.append([row.get(c) for c in columns])
        buf = io.BytesIO()
        book.save(buf)
        return {"format": "xlsx", "bytes_base64": base64.b64encode(buf.getvalue()).decode("ascii"),
                "filename": re.sub(r"[^A-Za-z0-9_-]+", "_", args.title)[:60] + ".xlsx"}

    # --------------------------------------------------------------- system
    def _offline(self) -> bool:
        try:
            from config import settings
            from backend.security.offline_policy import offline_mode
            return bool(offline_mode(settings, bool(settings.is_production)))
        except Exception:
            return False

    def _health(self, args, ctx):
        db_ok = False
        try:
            db_ok = bool(self.db and (self.db.execute_query("SELECT 1 AS ok") or {}).get("success"))
        except Exception:
            db_ok = False
        count = 0
        try:
            count = int(self.kb.collection.count()) if self.kb is not None else 0
        except Exception:
            count = 0
        return {"database": db_ok, "knowledge_base_examples": count, "providers": self.providers,
                "offline_mode": self._offline()}

    def _capabilities(self, args, ctx):
        return {"tools": self.names(), "version": VERSION, "offline_mode": self._offline(),
                "vector_store": str(getattr(self.config, "vector_store", "chroma") or "chroma")}

    # ------------------------------------------------------------- catalogue
    def _register_all(self):
        r = self._register
        r("database.search_schema", "Find tables and columns relevant to a question.", SearchSchemaIn, SearchSchemaOut, self._search_schema, 5)
        r("database.describe_table", "Columns, keys and description of one exposed table.", DescribeTableIn, DescribeTableOut, self._describe_table, 5)
        r("database.get_relationships", "How the exposed tables join, with the rules that matter.", Empty, RelationshipsOut, self._relationships, 5)
        r("database.list_allowed_tables", "The tables generated SQL may read.", Empty, AllowedTablesOut, self._allowed_tables, 5)
        r("vanna.retrieve_sql_examples", "Verified question -> SQL examples similar to a question.", RetrieveExamplesIn, RetrieveExamplesOut, self._retrieve_examples, 20)
        r("vanna.retrieve_ddl", "DDL of the exposed tables (all, or the ones named).", RetrieveDdlIn, RetrieveDdlOut, self._retrieve_ddl, 5)
        r("vanna.retrieve_business_context", "Business rules and join notes relevant to a question.", BusinessContextIn, BusinessContextOut, self._business_context, 5)
        r("sql.generate", "Generate SQL for a question through the deployment's generator.", GenerateIn, GenerateOut, self._generate, 180)
        r("sql.explain", "Plain description of what a statement reads and computes (no execution).", SqlIn, ExplainOut, self._explain, 5)
        r("sql.validate", "Run the AST policy: statement type, tables, scope, limits.", SqlIn, ValidateOut, self._validate, 10)
        r("security.authorize_query", "Authorization verdict for the caller: policy plus camera scope.", SqlIn, ValidateOut, self._authorize, 10)
        r("security.enforce_readonly", "Is this a single read-only SELECT?", SqlIn, ReadonlyOut, self._enforce_readonly, 5)
        r("security.check_allowed_tables", "Which referenced tables are outside the allow-list.", SqlIn, AllowedCheckOut, self._check_allowed_tables, 10)
        r("security.apply_limits", "The statement with the row limit enforced.", ApplyLimitsIn, ApplyLimitsOut, self._apply_limits, 10)
        r("database.execute_readonly", "Execute an authorized SELECT through the read-only role.", ExecuteIn, ExecuteOut, self._execute_readonly, 60, reads_data=True)
        r("database.explain_query", "Query plan (when the policy allows EXPLAIN).", SqlIn, ExplainQueryOut, self._explain_query, 30)
        r("database.cancel_query", "Cancel a running request by id.", CancelIn, CancelOut, self._cancel, 5)
        r("analytics.aggregate", "Count/sum/avg/min/max rows by a key.", AggregateIn, AggregateOut, self._aggregate, 10)
        r("analytics.compare", "Compare two groups on one value.", CompareIn, CompareOut, self._compare, 10)
        r("analytics.trend", "Order points on x and describe the change in y.", TrendIn, TrendOut, self._trend, 10)
        r("analytics.percentage_change", "Change and percentage change between two figures.", PctChangeIn, PctChangeOut, self._pct_change, 5)
        r("report.generate_csv", "CSV of rows (base64).", ReportIn, ReportOut, self._csv, 30)
        r("report.generate_pdf", "PDF report (base64) through the existing builder.", ReportIn, ReportOut, self._pdf, 60)
        r("report.generate_excel", "Excel workbook of rows (base64).", ReportIn, ReportOut, self._excel, 60)
        r("system.health", "Database, knowledge base and providers.", Empty, HealthOut, self._health, 10)
        r("system.capabilities", "The tool catalogue and deployment mode.", Empty, CapabilitiesOut, self._capabilities, 5)


TOOL_NAMES = sorted([
    "database.search_schema", "database.describe_table", "database.get_relationships", "database.list_allowed_tables",
    "vanna.retrieve_sql_examples", "vanna.retrieve_ddl", "vanna.retrieve_business_context",
    "sql.generate", "sql.explain", "sql.validate",
    "security.authorize_query", "security.enforce_readonly", "security.check_allowed_tables", "security.apply_limits",
    "database.execute_readonly", "database.explain_query", "database.cancel_query",
    "analytics.aggregate", "analytics.compare", "analytics.trend", "analytics.percentage_change",
    "report.generate_pdf", "report.generate_excel", "report.generate_csv",
    "system.health", "system.capabilities",
])
